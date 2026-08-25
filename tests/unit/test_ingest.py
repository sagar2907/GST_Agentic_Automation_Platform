"""Ingest tests, against workbooks written by a library that is not mine.

The reader in ``ingest.xlsx`` is hand-written, so testing it against fixtures
this repository also hand-writes would prove only that two pieces of my own
code agree about the format. Every workbook below is written by ``openpyxl``
-- a dev-only dependency, never imported by the package -- so the fixtures are
an independent implementation of the format's write side.

That is the whole reason it is a dependency. It is not used to read anything.
"""

from __future__ import annotations

import re
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from gst_recon.ingest import coerce, schema
from gst_recon.ingest.loader import (
    HEADER_ROWS,
    IngestRefusedError,
    load_purchase_register,
    read_table,
)
from gst_recon.ingest.xlsx import XlsxError, column_index, read_sheet, serial_to_date

HEADERS = [
    "Supplier GSTIN",
    "Invoice No",
    "Invoice Date",
    "Taxable Value",
    "CGST",
    "SGST",
    "IGST",
    "Reverse Charge",
]
GSTIN = "27AAAPA1234A1Z5"


def _write(path: Path, rows: list[list[object]], sheet_name: str = "Register") -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def _register(tmp_path: Path, body: list[list[object]], name: str = "register.xlsx") -> Path:
    return _write(tmp_path / name, [HEADERS, *body])


def _row(
    number: str = "INV/001",
    when: object = date(2026, 6, 14),
    *,
    taxable: object = 1000,
    cgst: object = 90,
    sgst: object = 90,
    igst: object = 0,
    rcm: str = "No",
    gstin: str = GSTIN,
) -> list[object]:
    return [gstin, number, when, taxable, cgst, sgst, igst, rcm]


# --------------------------------------------------------------------------
# The reader


def test_a_number_arrives_as_the_text_the_file_stores(tmp_path) -> None:
    """The point of not using a spreadsheet library.

    A currency cell must reach Decimal as the literal the file holds. Anything
    that goes via float has already made a decision about the value before this
    system sees it.
    """
    path = _write(tmp_path / "amounts.xlsx", [["amount"], [Decimal("1234.56")]])
    cell = read_sheet(path).body[0][0]
    assert cell == "1234.56"
    assert Decimal(cell) == Decimal("1234.56")


def test_a_date_cell_is_resolved_from_its_number_format(tmp_path) -> None:
    """A date is a serial number; only the applied format says it is a date."""
    path = _write(tmp_path / "dates.xlsx", [["when"], [date(2026, 6, 14)]])
    assert read_sheet(path).body[0][0] == "2026-06-14"


def test_a_quantity_is_not_mistaken_for_a_date(tmp_path) -> None:
    path = _write(tmp_path / "mixed.xlsx", [["when", "qty"], [date(2026, 6, 14), 45000]])
    when, quantity = read_sheet(path).body[0]
    assert when == "2026-06-14"
    assert quantity == "45000"


def test_a_gap_in_a_row_keeps_the_columns_aligned(tmp_path) -> None:
    """Sparse rows omit empty cells, so position comes from the reference."""
    path = _write(tmp_path / "sparse.xlsx", [["a", "b", "c"], ["left", None, "right"]])
    assert read_sheet(path).body[0] == ["left", "", "right"]


def test_trailing_blank_rows_go_and_interior_ones_stay(tmp_path) -> None:
    """A trailing blank row is how the file was saved. An interior one is data.

    Dropping the interior row would change the row numbers every rejection
    message quotes, sending a reviewer to the wrong line.
    """
    path = _write(tmp_path / "gaps.xlsx", [["a"], ["first"], [None], ["third"], [None], [None]])
    rows = read_sheet(path).body
    assert rows == [["first"], [], ["third"]]


def test_a_named_sheet_can_be_chosen_and_a_missing_one_is_named(tmp_path) -> None:
    path = tmp_path / "two.xlsx"
    book = Workbook()
    book.active.title = "First"
    book.active.append(["a"])
    book.create_sheet("Second").append(["b"])
    book.save(path)

    assert read_sheet(path, "Second").rows == [["b"]]
    with pytest.raises(XlsxError, match="First, Second"):
        read_sheet(path, "Third")


def test_something_that_is_not_a_workbook_is_refused(tmp_path) -> None:
    path = tmp_path / "not-really.xlsx"
    path.write_text("supplier,invoice\n", encoding="utf-8")
    with pytest.raises(XlsxError, match="not a zip archive"):
        read_sheet(path)


def test_a_zip_that_is_not_a_workbook_is_refused(tmp_path) -> None:
    path = tmp_path / "empty.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("hello.txt", "nothing here")
    with pytest.raises(XlsxError):
        read_sheet(path)


@pytest.mark.parametrize(
    ("reference", "expected"),
    [("A1", 0), ("B2", 1), ("Z9", 25), ("AA1", 26), ("AB7", 27), ("BA1", 52)],
)
def test_column_references_decode(reference: str, expected: int) -> None:
    assert column_index(reference) == expected


def test_the_1900_leap_year_bug_is_accounted_for() -> None:
    """Excel believes 29 February 1900 existed. Serials either side prove it.

    Without the correction every date in a workbook is off by one, which is a
    silent error: 31 March and 1 April are both plausible invoice dates and one
    of them is in the next return period.
    """
    assert serial_to_date(59) == date(1900, 2, 28)
    assert serial_to_date(61) == date(1900, 3, 1)
    assert serial_to_date(45822) == date(2025, 6, 14)
    with pytest.raises(XlsxError, match="non-existent"):
        serial_to_date(60)


# --------------------------------------------------------------------------
# Header mapping


@pytest.mark.parametrize(
    "heading",
    ["GSTIN", "Supplier GSTIN", "GSTIN of Supplier", "gstin_no", "Vendor GSTIN", " party  gstin "],
)
def test_supplier_gstin_is_found_under_the_names_it_arrives_with(heading: str) -> None:
    mapped = schema.map_headers([heading, "Invoice No", "Invoice Date", "Taxable Value", "CGST"])
    assert mapped.usable, mapped.problems
    assert mapped.index("supplier_gstin") == 0


@pytest.mark.parametrize(
    "heading", ["Recipient GSTIN", "Our GSTIN", "Buyer GSTIN", "GSTIN of Recipient"]
)
def test_the_taxpayers_own_gstin_is_never_read_as_the_suppliers(heading: str) -> None:
    """The failure this table exists to prevent.

    A loose alias would see "GSTIN" inside "Recipient GSTIN" and map the
    taxpayer's own registration as the supplier on every row, producing a
    register that reconciles against nothing and blames the wrong party.
    """
    mapped = schema.map_headers([heading, "Invoice No", "Invoice Date", "Taxable Value", "CGST"])
    assert not mapped.usable
    assert any("supplier_gstin" in problem for problem in mapped.problems)
    assert mapped.ignored[0] == "recipient identity"


def test_two_columns_meaning_the_same_field_is_a_refusal() -> None:
    """Not a tie-break. The file draws a distinction this code cannot read."""
    mapped = schema.map_headers(
        ["Supplier GSTIN", "Party GSTIN", "Invoice No", "Invoice Date", "Taxable Value", "CGST"]
    )
    assert not mapped.usable
    problem = " ".join(mapped.problems)
    assert "'Supplier GSTIN'" in problem
    assert "'Party GSTIN'" in problem


def test_a_missing_required_column_names_what_was_actually_seen() -> None:
    """The message is read by someone with the spreadsheet open."""
    mapped = schema.map_headers(["Supplier GSTIN", "Invoice No", "Taxable Value", "CGST"])
    assert not mapped.usable
    assert any("invoice_date" in problem for problem in mapped.problems)
    assert any("'Supplier GSTIN'" in problem for problem in mapped.problems)


def test_a_register_with_no_tax_column_is_refused() -> None:
    mapped = schema.map_headers(["Supplier GSTIN", "Invoice No", "Invoice Date", "Taxable Value"])
    assert not mapped.usable
    assert any("no tax column" in problem for problem in mapped.problems)


def test_a_derived_total_is_not_mistaken_for_a_tax_head() -> None:
    mapped = schema.map_headers(
        ["Supplier GSTIN", "Invoice No", "Invoice Date", "Taxable Value", "Total Tax"]
    )
    assert not mapped.usable
    assert any("no tax column" in problem for problem in mapped.problems)


def test_unrecognised_columns_are_reported_rather_than_silently_dropped() -> None:
    mapped = schema.map_headers([*HEADERS, "Narration", "Cost Centre"])
    assert mapped.usable, mapped.problems
    assert sorted(mapped.unrecognised.values()) == ["Cost Centre", "Narration"]


# --------------------------------------------------------------------------
# Coercion


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("1234.56", "1234.56"),
        ("1,23,456.78", "123456.78"),
        ("₹ 1,234.56", "1234.56"),
        ("INR 1234.56", "1234.56"),
        ("(1,234.56)", "-1234.56"),
        ("1234.56-", "-1234.56"),
        ("1234.56 Cr", "1234.56"),
        ("  1234.5  ", "1234.50"),
        ("0", "0.00"),
    ],
)
def test_amounts_arrive_as_the_figure_a_person_would_read(text: str, expected: str) -> None:
    assert coerce.parse_amount(text) == Decimal(expected)


@pytest.mark.parametrize("text", ["", "  ", "1234 approx", "n/a", "-", "12.34.56", "abc"])
def test_an_amount_that_is_not_one_is_refused(text: str) -> None:
    with pytest.raises(coerce.CoercionError):
        coerce.parse_amount(text)


def test_amounts_are_exact_where_floats_would_not_be() -> None:
    total = coerce.parse_amount("0.1") + coerce.parse_amount("0.2")
    assert total == coerce.parse_amount("0.3")


def test_one_unambiguous_row_settles_the_whole_columns_convention() -> None:
    column = coerce.infer_date_order(["01/02/2026", "14/02/2026", "03/03/2026"])
    assert column.order is coerce.DateOrder.DAY_FIRST
    assert column.evidence == "14/02/2026"

    column = coerce.infer_date_order(["01/02/2026", "02/14/2026"])
    assert column.order is coerce.DateOrder.MONTH_FIRST
    assert column.evidence == "02/14/2026"


def test_a_column_that_contradicts_itself_resolves_to_nothing() -> None:
    """Both readings are refuted by some row, so the column has no convention."""
    column = coerce.infer_date_order(["14/02/2026", "02/14/2026"])
    assert column.order is coerce.DateOrder.UNRESOLVED
    assert "14/02/2026" in column.evidence
    assert "02/14/2026" in column.evidence


def test_a_column_where_every_day_falls_before_the_13th_stays_unresolved() -> None:
    column = coerce.infer_date_order(["01/02/2026", "03/04/2026", "05/06/2026"])
    assert column.order is coerce.DateOrder.UNRESOLVED
    assert column.evidence == ""


def test_iso_values_carry_no_evidence_about_the_numeric_ones() -> None:
    assert (
        coerce.infer_date_order(["2026-02-14", "2026-03-01"]).order is coerce.DateOrder.UNRESOLVED
    )


def test_the_same_text_is_two_different_dates_under_the_two_conventions() -> None:
    assert coerce.parse_date("03/04/2026", coerce.DateOrder.DAY_FIRST) == date(2026, 4, 3)
    assert coerce.parse_date("03/04/2026", coerce.DateOrder.MONTH_FIRST) == date(2026, 3, 4)


def test_an_iso_date_ignores_the_columns_convention() -> None:
    for order in (coerce.DateOrder.DAY_FIRST, coerce.DateOrder.MONTH_FIRST):
        assert coerce.parse_date("2026-04-03", order) == date(2026, 4, 3)


def test_an_unresolved_column_refuses_rather_than_picking_a_locale() -> None:
    with pytest.raises(coerce.CoercionError, match="ambiguous"):
        coerce.parse_date("03/04/2026", coerce.DateOrder.UNRESOLVED)


def test_a_date_that_does_not_exist_is_refused() -> None:
    with pytest.raises(coerce.CoercionError, match="not a real date"):
        coerce.parse_date("31/02/2026", coerce.DateOrder.DAY_FIRST)


@pytest.mark.parametrize("text", ["UNKNOWN", "maybe", "2", "-"])
def test_a_flag_that_is_neither_yes_nor_no_is_refused(text: str) -> None:
    """Reading UNKNOWN as No moves the tax liability to the wrong party."""
    with pytest.raises(coerce.CoercionError):
        coerce.parse_flag(text)


def test_an_identifier_is_tidied_but_never_altered() -> None:
    assert coerce.normalise_identifier(" 27aaapa1234a1z5 ") == GSTIN
    assert coerce.normalise_identifier("inv / 001") == "INV/001"


# --------------------------------------------------------------------------
# The loader


def test_a_clean_register_loads(tmp_path) -> None:
    path = _register(
        tmp_path,
        [_row(), _row("INV/002", date(2026, 6, 15), taxable=2000, cgst=0, sgst=0, igst=360)],
    )
    report = load_purchase_register(path)

    assert len(report.lines) == 2
    assert report.rejected == []
    first, second = report.lines
    assert first.supplier_gstin == GSTIN
    assert first.invoice_number == "INV/001"
    assert first.invoice_date == date(2026, 6, 14)
    assert first.tax.taxable == Decimal("1000.00")
    assert first.tax.total_tax == Decimal("180.00")
    assert second.tax.is_interstate


def test_a_row_number_points_at_the_line_a_person_sees(tmp_path) -> None:
    """Off by one here sends a reviewer to the wrong line of a long register."""
    path = _register(tmp_path, [_row(), _row("INV/002", taxable="not a number")])
    report = load_purchase_register(path)

    assert len(report.lines) == 1
    assert [rejection.row for rejection in report.rejected] == [HEADER_ROWS + 2]
    assert report.lines[0].source_row == HEADER_ROWS + 1


def test_a_blank_row_does_not_shift_the_rows_after_it(tmp_path) -> None:
    """The reason interior blanks are preserved rather than closed up.

    A register with a spacer row between sections is ordinary. If reading it
    dropped that row, every row number after it would be one too low, and the
    rejection message for row 400 would send a reviewer to row 399.
    """
    body: list[list[object]] = [_row("INV/001"), [None], _row("INV/002", taxable="n/a")]
    report = load_purchase_register(_register(tmp_path, body))

    assert [line.source_row for line in report.lines] == [HEADER_ROWS + 1]
    assert [rejection.row for rejection in report.rejected] == [HEADER_ROWS + 3]


def test_one_bad_row_does_not_cost_the_client_the_register(tmp_path) -> None:
    body = [_row(f"INV/{index:03d}") for index in range(5)]
    body[2] = _row("INV/002", taxable="")
    report = load_purchase_register(_register(tmp_path, body))

    assert len(report.lines) == 4
    assert len(report.rejected) == 1
    assert report.rejection_rate == pytest.approx(0.2)


def test_a_row_is_rejected_whole_rather_than_patched(tmp_path) -> None:
    """A default in place of a failed field is a number nobody supplied."""
    path = _register(tmp_path, [_row(), _row("INV/002", rcm="UNKNOWN")])
    report = load_purchase_register(path)

    assert [line.invoice_number for line in report.lines] == ["INV/001"]
    assert "neither yes nor no" in report.rejected[0].reason


def test_a_document_carrying_both_tax_regimes_is_rejected(tmp_path) -> None:
    """No single supply is both intra-state and inter-state.

    A row with all three heads populated is either two supplies on one line or
    the wrong columns, and both are things to stop for.
    """
    path = _register(tmp_path, [_row(), _row("INV/002", cgst=90, sgst=90, igst=180)])
    report = load_purchase_register(path)

    assert len(report.lines) == 1
    assert "both intra-state and inter-state" in report.rejected[0].reason


def test_a_file_whose_header_cannot_be_mapped_loads_nothing(tmp_path) -> None:
    path = _write(tmp_path / "wrong.xlsx", [["Alpha", "Beta", "Gamma"], [1, 2, 3]])
    with pytest.raises(IngestRefusedError, match="cannot be mapped"):
        load_purchase_register(path)


def test_a_wholly_ambiguous_date_column_stops_the_file(tmp_path) -> None:
    """One file problem, not forty row problems.

    Every day in this register falls on or before the 12th, so nothing in the
    file says which way round its dates are written. Guessing would move
    documents across a period boundary silently.
    """
    body = [_row("INV/001", "01/02/2026"), _row("INV/002", "03/04/2026")]
    with pytest.raises(IngestRefusedError, match="cannot be read"):
        load_purchase_register(_register(tmp_path, body))


def test_the_caller_may_assert_the_date_order_the_file_cannot(tmp_path) -> None:
    """An assertion by someone who knows the file's provenance is not a guess."""
    body = [_row("INV/001", "01/02/2026"), _row("INV/002", "03/04/2026")]
    path = _register(tmp_path, body)

    report = load_purchase_register(path, date_order=coerce.DateOrder.DAY_FIRST)
    assert [line.invoice_date for line in report.lines] == [date(2026, 2, 1), date(2026, 4, 3)]

    report = load_purchase_register(path, date_order=coerce.DateOrder.MONTH_FIRST)
    assert [line.invoice_date for line in report.lines] == [date(2026, 1, 2), date(2026, 3, 4)]


def test_an_all_iso_column_needs_no_convention(tmp_path) -> None:
    body = [_row("INV/001", "2026-02-01"), _row("INV/002", "2026-04-03")]
    report = load_purchase_register(_register(tmp_path, body))
    assert [line.invoice_date for line in report.lines] == [date(2026, 2, 1), date(2026, 4, 3)]


def test_every_row_failing_is_reported_as_one_file_problem(tmp_path) -> None:
    body = [_row(f"INV/{index:03d}", taxable="n/a") for index in range(4)]
    with pytest.raises(IngestRefusedError, match="every row was rejected"):
        load_purchase_register(_register(tmp_path, body))


def test_a_header_and_no_rows_is_refused(tmp_path) -> None:
    with pytest.raises(IngestRefusedError, match="no rows"):
        load_purchase_register(_register(tmp_path, []))


def test_a_csv_exported_from_excel_loads_despite_its_byte_order_mark(tmp_path) -> None:
    """The BOM would otherwise become part of the first heading."""
    path = tmp_path / "register.csv"
    lines = [",".join(HEADERS), f"{GSTIN},INV/001,2026-06-14,1000,90,90,0,No"]
    path.write_text("\n".join(lines), encoding="utf-8-sig")

    assert read_table(path)[0][0] == "Supplier GSTIN"
    report = load_purchase_register(path)
    assert len(report.lines) == 1
    assert report.lines[0].tax.taxable == Decimal("1000.00")


def test_an_unsupported_container_is_refused(tmp_path) -> None:
    path = tmp_path / "register.pdf"
    path.write_bytes(b"%PDF-1.4")
    with pytest.raises(IngestRefusedError, match=re.escape("expected .xlsx or .csv")):
        load_purchase_register(path)


def test_the_summary_says_how_the_file_was_interpreted(tmp_path) -> None:
    """A caller that never learns it dropped rows reconciles a partial register.

    An inward document missing at the cut-off is treated as accepted, so a
    silent drop is not a reporting nicety.
    """
    body = [_row("INV/001", "14/06/2026"), _row("INV/002", "15/06/2026", taxable="n/a")]
    report = load_purchase_register(_register(tmp_path, body))
    summary = report.summary()

    assert "1 of 2 rows loaded" in summary
    assert "1 rejected" in summary
    assert "day_first" in summary
